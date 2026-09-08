"""简历文件解析：PDF / DOCX / 纯文本 → 纯文本。"""

import io
from pathlib import Path

MAX_TEXT_LEN = 30000

ALLOWED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB


class ParseError(Exception):
    """文件解析失败，message 面向用户展示。"""


def _truncate(text: str) -> str:
    text = text.strip()
    if len(text) > MAX_TEXT_LEN:
        text = text[:MAX_TEXT_LEN] + "\n…（内容过长，已截断）"
    return text


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8", "gbk", "utf-16", "big5"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, ValueError):
            continue
    return data.decode("utf-8", errors="ignore")


def _parse_pdf(data: bytes) -> str:
    import pdfplumber

    pages: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    text = "\n".join(pages).strip()
    if not text:
        raise ParseError("无法从 PDF 中提取到文本，可能是扫描件/图片型简历，请粘贴文本后重试。")
    return text


def _parse_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:  # python-docx 对损坏文件抛多种异常
        raise ParseError("Word 文件解析失败，文件可能已损坏。") from e

    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # 很多简历模板用表格排版
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(c for c in cells if c))
    text = "\n".join(p for p in parts if p and p.strip()).strip()
    if not text:
        raise ParseError("无法从 Word 文件中提取到文本内容。")
    return text


def parse_resume_file(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise ParseError("不支持的文件类型，仅支持 PDF / DOCX / TXT / MD。")
    if len(data) == 0:
        raise ParseError("文件内容为空。")
    if len(data) > MAX_FILE_SIZE:
        raise ParseError("文件超过 10MB 限制。")

    if suffix == ".pdf":
        text = _parse_pdf(data)
    elif suffix == ".docx":
        text = _parse_docx(data)
    else:
        text = _decode_text(data)

    if not text.strip():
        raise ParseError("文件解析结果为空，请检查文件内容。")
    return _truncate(text)
